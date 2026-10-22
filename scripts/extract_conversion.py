#!/usr/bin/env python3
"""
从「日维度_搜索方式转化情况」Excel 中提取指定搜索方式的关键指标，
生成带格式的汇总 Excel。

用法:
    python3 extract_conversion.py <excel文件路径> [起始日期] [结束日期]

日期格式: YYYYMMDD (如 20260907)
不指定日期时自动提取全部日期。
"""

import sys
import os
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

# 固定提取的搜索方式及顺序
TARGET_METHODS = [
    "手动文字-AI搜",
    "底纹词-AI搜",
    "语音输入-AI搜",
    "试试这样问-AI搜",
    "你最近问过-AI搜",
    "推荐词-前置页-AI搜",
]

# 需要提取的字段名（与源 Excel 表头精确匹配，列位置可变）
FIELD_NAMES = ["搜索pv", "搜索uv", "总提问条数", "直接订单数"]

OUTPUT_HEADERS = ["日期", "搜索方式"] + FIELD_NAMES


def build_col_index(ws):
    """扫描表头行，返回 {列名: 1-based列索引}，兼容带前导空格的列名。"""
    col_map = {}
    for col in range(1, ws.max_column + 1):
        val = ws.cell(1, col).value
        if val is not None:
            col_map[str(val).strip()] = col
    return col_map


def extract_data(src_path, date_start=None, date_end=None):
    """从源 Excel 读取数据，返回 [(date, method, pv, uv, tqs, ordn), ...]"""
    wb = openpyxl.load_workbook(src_path, data_only=True)
    ws = wb["日维度_搜索方式转化情况"]

    col_map = build_col_index(ws)

    # 检查所有必需列是否存在
    missing = [f for f in FIELD_NAMES if f not in col_map]
    if missing:
        print("错误: 源 Excel 中缺少以下列: %s" % ", ".join(missing))
        print("可用列: %s" % ", ".join(sorted(col_map.keys())))
        sys.exit(1)

    field_indices = [col_map[f] for f in FIELD_NAMES]

    target_set = set(TARGET_METHODS)
    rows = []
    last_date = None

    for r in range(2, ws.max_row + 1):
        d = ws.cell(r, 1).value
        if d is not None:
            last_date = str(d)
        method = ws.cell(r, 3).value
        if method not in target_set:
            continue

        # 日期过滤
        if date_start and last_date < date_start:
            continue
        if date_end and last_date > date_end:
            continue

        vals = []
        for col_idx in field_indices:
            v = ws.cell(r, col_idx).value
            vals.append(v if v is not None else 0)
        rows.append((last_date, method, *vals))

    return rows


def write_output(rows, out_path):
    """写入带格式的 Excel"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "搜索方式转化情况"

    # 样式
    header_font = Font(bold=True, size=11)
    header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    center = Alignment(horizontal="center", vertical="center")
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    # 表头
    for col, h in enumerate(OUTPUT_HEADERS, 1):
        cell = ws.cell(1, col, h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = border

    # 数据
    for i, row_data in enumerate(rows, 2):
        for col, v in enumerate(row_data, 1):
            cell = ws.cell(i, col, v)
            cell.alignment = center
            cell.border = border

    # 列宽
    widths = [12, 18, 12, 12, 12, 12]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    wb.save(out_path)


def main():
    if len(sys.argv) < 2:
        print("用法: python3 extract_conversion.py <excel文件路径> [起始日期] [结束日期]")
        print("日期格式: YYYYMMDD")
        sys.exit(1)

    src_path = sys.argv[1]
    date_start = sys.argv[2] if len(sys.argv) > 2 else None
    date_end = sys.argv[3] if len(sys.argv) > 3 else None

    if not os.path.exists(src_path):
        print("错误: 文件不存在 - %s" % src_path)
        sys.exit(1)

    rows = extract_data(src_path, date_start, date_end)
    if not rows:
        print("错误: 未找到匹配数据，请检查日期范围和搜索方式")
        sys.exit(1)

    # 排序: 日期升序 + 搜索方式固定顺序
    method_order = {m: i for i, m in enumerate(TARGET_METHODS)}
    rows.sort(key=lambda r: (r[0], method_order.get(r[1], 99)))

    # 生成输出路径
    dates = [r[0] for r in rows]
    d_min, d_max = min(dates), max(dates)
    if date_start or date_end:
        name_range = "%s-%s" % (date_start or d_min, date_end or d_max)
    else:
        name_range = "%s-%s" % (d_min, d_max)

    # 输出到 outputs 目录
    cwd = os.getcwd()
    out_dir = os.path.join(cwd, "outputs")
    os.makedirs(out_dir, exist_ok=True)
    out_path = os.path.join(out_dir, "搜索方式转化情况_%s.xlsx" % name_range)

    write_output(rows, out_path)

    print("已生成: %s" % out_path)
    print("共 %d 行数据 (%d 天 × %d 种搜索方式)" % (len(rows), len(set(dates)), len(TARGET_METHODS)))
    print("日期范围: %s ~ %s" % (d_min, d_max))


if __name__ == "__main__":
    main()

